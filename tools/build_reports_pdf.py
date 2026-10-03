#!/usr/bin/env python3
"""Fallback PDF export of the existing report DOCX when LibreOffice is unavailable.

Requires python-docx and ReportLab. Fonts are supplied explicitly; no system-font
substitution or source modification is performed. This is a bounded renderer for
tools/build_reports.py output, rather than a general Word layout implementation.
"""
import argparse
from io import BytesIO
from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph as WordParagraph
from docx.text.run import Run
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, PageBreak, PageTemplate, Paragraph, Spacer,
    Table as PDFTable, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

PAGE_WIDTH, PAGE_HEIGHT = A4
TEXT_WIDTH = PAGE_WIDTH - 45 * mm
TEXT_HEIGHT = PAGE_HEIGHT - 40 * mm


def inherited(paragraph, attr, font=False):
    value = getattr(paragraph.style.font if font else paragraph.paragraph_format, attr)
    if not font and value is not None:
        return value
    style = paragraph.style
    while style is not None:
        value = getattr(style.font if font else style.paragraph_format, attr)
        if value is not None:
            return value
        style = style.base_style
    return None


def register_fonts(directory):
    names = {}
    for family, token in (("TNR", "TimesNewRoman"), ("Courier", "CourierNew")):
        for suffix, marker in (("", "PSMT"), ("-Bold", "PS-BoldMT"),
                               ("-Italic", "PS-ItalicMT")):
            matches = sorted(p for p in directory.glob("*.ttf")
                             if token in p.name and marker in p.name)
            if len(matches) != 1:
                raise ValueError(f"Need exactly one {family}{suffix} font in {directory}")
            name = family + suffix
            font = TTFont(name, str(matches[0]))
            pdfmetrics.registerFont(font)
            names[name] = font
        pdfmetrics.registerFontFamily(family, normal=family, bold=family + "-Bold",
                                      italic=family + "-Italic", boldItalic=family + "-Bold")
    return names


class ReportPDF(BaseDocTemplate):
    def __init__(self, output, source):
        super().__init__(output, pagesize=A4, leftMargin=30 * mm, rightMargin=15 * mm,
                         topMargin=20 * mm, bottomMargin=20 * mm,
                         title=source.core_properties.title,
                         author=source.core_properties.author,
                         subject=source.core_properties.subject)
        self.first_footer = " ".join(p.text for p in source.sections[0].first_page_footer.paragraphs)
        frame = Frame(self.leftMargin, self.bottomMargin, TEXT_WIDTH, TEXT_HEIGHT,
                      leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates(PageTemplate(id="report", frames=frame, onPage=self.footer))

    def footer(self, canvas, doc):
        canvas.saveState()
        canvas.setFont("TNR", 14)
        canvas.drawCentredString(PAGE_WIDTH / 2, 10 * mm,
                                 self.first_footer if doc.page == 1 else str(doc.page))
        canvas.restoreState()

    def afterFlowable(self, flowable):
        if hasattr(flowable, "toc_level"):
            self.notify("TOCEntry", (flowable.toc_level, flowable.getPlainText(), self.page))


class Renderer:
    def __init__(self, document, fonts, symbol_font=None):
        self.document = document
        self.fonts = fonts
        self.symbol_font = symbol_font
        self.symbol_fallbacks = set()
        self.heading_count = sum(p.style.name.startswith("Heading ") for p in document.paragraphs)

    def run_properties(self, run, paragraph):
        family = "Courier" if "Courier" in (run.font.name or inherited(paragraph, "name", True) or "") else "TNR"
        bold = run.bold if run.bold is not None else inherited(paragraph, "bold", True)
        italic = run.italic if run.italic is not None else inherited(paragraph, "italic", True)
        if bold and italic and run.text.strip():
            raise ValueError("A bold-italic run needs a supplied bold-italic font; refusing to drop its formatting")
        name = family + ("-Bold" if bold else "-Italic" if italic else "")
        size = run.font.size or inherited(paragraph, "size", True)
        return name, size.pt if size is not None else 14

    def markup(self, text, name, size):
        missing = sorted({c for c in text if not c.isspace() and ord(c) not in self.fonts[name].face.charToGlyph})
        if missing:
            symbols = self.fonts.get("Symbols")
            unsupported = [c for c in missing if c.isalnum() or symbols is None
                           or ord(c) not in symbols.face.charToGlyph]
            if unsupported:
                raise ValueError(f"Font {name} lacks required glyphs: {''.join(unsupported)!r}")
            self.symbol_fallbacks.update((c, name) for c in missing)
            pieces, chunk, current_font = [], [], None
            for char in text:
                font = "Symbols" if char in missing else name
                if current_font is not None and font != current_font:
                    pieces.append(self.markup("".join(chunk), current_font, size))
                    chunk = []
                chunk.append(char)
                current_font = font
            if chunk:
                pieces.append(self.markup("".join(chunk), current_font, size))
            return "".join(pieces)
        text = escape(text).replace("\n", "<br/>").replace("\t", "&#160;" * 4)
        # Preserve multiple spaces in code and signature lines.
        text = re.sub(r" {2,}", lambda m: "&#160;" * len(m.group()), text)
        return f'<font name="{name}" size="{size:g}">{text}</font>'

    def style(self, paragraph, in_table=False):
        size = inherited(paragraph, "size", True)
        size = size.pt if size is not None else 14
        sizes = [(r.font.size.pt if r.font.size is not None else size)
                 for r in paragraph.runs if r.text.strip()]
        if sizes:
            size = max(sizes)
        spacing = inherited(paragraph, "line_spacing") or 1.5
        leading = spacing.pt if hasattr(spacing, "pt") else size * float(spacing)
        def points(attr):
            value = inherited(paragraph, attr)
            return value.pt if value is not None else 0
        alignment = inherited(paragraph, "alignment")
        alignment = {0: TA_LEFT, 1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}.get(alignment, TA_LEFT)
        return ParagraphStyle("docx", fontName="TNR", fontSize=size, leading=leading,
                              alignment=alignment, firstLineIndent=points("first_line_indent"),
                              leftIndent=points("left_indent"), rightIndent=points("right_indent"),
                              spaceBefore=points("space_before"), spaceAfter=points("space_after"),
                              keepWithNext=bool(inherited(paragraph, "keep_with_next")) and not in_table,
                              splitLongWords=True, allowWidows=0, allowOrphans=0)

    def paragraph(self, paragraph, in_table=False):
        if any("TOC " in (node.text or "") for node in paragraph._p.xpath(".//w:instrText")):
            toc = TableOfContents()
            spacing = 1.5 if self.heading_count <= 24 else 1.15 if self.heading_count <= 34 else 1
            toc.levelStyles = [ParagraphStyle(f"toc{i}", fontName="TNR", fontSize=14,
                               leading=14 * spacing, leftIndent=i * 7.5 * mm,
                               firstLineIndent=0, spaceBefore=0, spaceAfter=0,
                               splitLongWords=True) for i in range(3)]
            toc.dotsMinLevel = 0
            return [toc]
        result, parts = [], []
        style = self.style(paragraph, in_table)
        def flush():
            if parts:
                item = Paragraph("".join(parts), style)
                if not in_table and re.fullmatch(r"Heading [123]", paragraph.style.name):
                    item.toc_level = int(paragraph.style.name[-1]) - 1
                result.append(item)
                parts.clear()
        if inherited(paragraph, "page_break_before") and not in_table:
            result.append(PageBreak())
        for element in paragraph._p.xpath(".//w:r"):
            run = Run(element, paragraph)
            name, size = self.run_properties(run, paragraph)
            for child in element:
                tag = child.tag
                if tag == qn("w:t"):
                    parts.append(self.markup(child.text or "", name, size))
                elif tag == qn("w:tab"):
                    parts.append("&#160;" * 4)
                elif tag == qn("w:br"):
                    if child.get(qn("w:type")) == "page":
                        flush()
                        result.append(PageBreak())
                    else:
                        parts.append("<br/>")
                elif tag == qn("w:drawing"):
                    flush()
                    blips = child.xpath(".//a:blip")
                    extents = child.xpath(".//wp:extent")
                    if len(blips) != 1 or not extents:
                        raise ValueError("Unsupported drawing: expected one inline image")
                    blob = self.document.part.related_parts[blips[0].get(qn("r:embed"))].blob
                    image = Image(BytesIO(blob))
                    width = int(extents[0].get("cx")) / 12700
                    height = int(extents[0].get("cy")) / 12700
                    scale = min(1, TEXT_WIDTH / width, (TEXT_HEIGHT - 50) / height)
                    image.drawWidth, image.drawHeight = width * scale, height * scale
                    image.hAlign = "CENTER"
                    image.spaceBefore, image.spaceAfter = style.spaceBefore, style.spaceAfter
                    image.keepWithNext = style.keepWithNext
                    result.append(image)
        flush()
        if not result:
            # Empty signature lines have real height; empty post-table spacers do too.
            result.append(Spacer(1, style.leading + style.spaceBefore + style.spaceAfter))
        return result

    def table(self, table):
        widths = [int(col.get(qn("w:w"))) / 20 for col in table._tbl.tblGrid]
        if not widths or sum(widths) <= 0:
            raise ValueError("Table has no usable OOXML column grid")
        if sum(widths) > TEXT_WIDTH:
            widths = [w * TEXT_WIDTH / sum(widths) for w in widths]
        data, spans, active_merges = [], [], {}
        repeat = 0
        for ri, row in enumerate(table._tbl.tr_lst):
            header = row.xpath("./w:trPr/w:tblHeader")
            if header and repeat == ri:
                repeat += 1
            values, ci = [""] * len(widths), 0
            for tc in row.tc_lst:
                cell = _Cell(tc, table)
                span = tc.grid_span
                merge = tc.xpath("./w:tcPr/w:vMerge")
                continuation = merge and merge[0].get(qn("w:val")) != "restart"
                if continuation:
                    if ci not in active_merges:
                        raise ValueError("Vertical table merge has no starting cell")
                    first, end = active_merges[ci]
                    spans[first] = ("SPAN", (ci, end), (ci + span - 1, ri))
                else:
                    values[ci] = [item for p in cell.paragraphs for item in self.paragraph(p, True)]
                    if merge or span > 1:
                        spans.append(("SPAN", (ci, ri), (ci + span - 1, ri)))
                    if merge:
                        active_merges[ci] = (len(spans) - 1, ri)
                    else:
                        active_merges.pop(ci, None)
                ci += span
            data.append(values)
        commands = [("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4.25),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4.25),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2)] + spans
        borders = table._tbl.xpath("./w:tblPr/w:tblBorders")
        if borders:
            edges = {node.tag.rsplit("}", 1)[-1]: node for node in borders[0]}
            border = next(iter(edges.values()))
            color = border.get(qn("w:color"), "000000")
            color = colors.HexColor("#" + color) if color != "auto" else colors.black
            weight = int(border.get(qn("w:sz"), "4")) / 8
            commands.append(("GRID" if "insideH" in edges else "BOX", (0, 0), (-1, -1), weight, color))
        for ri, row in enumerate(table._tbl.tr_lst):
            ci = 0
            for tc in row.tc_lst:
                shading = tc.xpath("./w:tcPr/w:shd")
                if shading and shading[0].get(qn("w:fill")) not in (None, "auto"):
                    commands.append(("BACKGROUND", (ci, ri), (ci + tc.grid_span - 1, ri),
                                     colors.HexColor("#" + shading[0].get(qn("w:fill")))))
                ci += tc.grid_span
        result = PDFTable(data, colWidths=widths, repeatRows=repeat,
                          splitByRow=1, splitInRow=1, hAlign="CENTER")
        result.setStyle(TableStyle(commands))
        return result

    def story(self):
        result = []
        for element in self.document.element.body:
            if element.tag == qn("w:p"):
                result.extend(self.paragraph(WordParagraph(element, self.document)))
            elif element.tag == qn("w:tbl"):
                result.append(self.table(Table(element, self.document)))
        return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--font-dir", type=Path, required=True)
    parser.add_argument("--symbol-font", type=Path,
                        help="Explicit font for missing non-letter symbols only; fallback use is reported")
    args = parser.parse_args()
    if args.pdf.exists():
        parser.error(f"Refusing to overwrite existing output: {args.pdf}")
    fonts = register_fonts(args.font_dir)
    if args.symbol_font:
        fonts["Symbols"] = TTFont("Symbols", str(args.symbol_font))
        pdfmetrics.registerFont(fonts["Symbols"])
    source = Document(args.docx)
    renderer = Renderer(source, fonts, args.symbol_font)
    story = renderer.story()  # Check fonts and input before reserving an output.
    if renderer.symbol_fallbacks:
        substitutions = ", ".join(f"{char!r} ({name})" for char, name in sorted(renderer.symbol_fallbacks))
        print(f"Explicit symbol fallback {args.symbol_font}: {substitutions}")
    args.pdf.parent.mkdir(parents=True, exist_ok=True)
    with args.pdf.open("xb") as output:
        ReportPDF(output, source).multiBuild(story)
    print(args.pdf)


if __name__ == "__main__":
    main()
